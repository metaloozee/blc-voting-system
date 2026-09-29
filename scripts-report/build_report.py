#!/usr/bin/env python3
"""Build Blockchain Mini Project Report DOCX matching SLRTCE college template.

Usage (from repo root):
  .venv-report/bin/python scripts-report/build_report.py
  libreoffice --headless --convert-to pdf --outdir report/ report/Blockchain_Mini_Project_Report.docx
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, Cm, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "report"
FIGS = REPORT / "figures"
OUT_DOCX = REPORT / "Blockchain_Mini_Project_Report.docx"

FONT = "Liberation Serif"
FONT_MONO = "Liberation Mono"


def set_run_font(run, size_pt=12, bold=False, italic=False, font=FONT):
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size_pt)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_para(
    doc,
    text="",
    *,
    size=12,
    bold=False,
    italic=False,
    align="left",
    space_before=0,
    space_after=6,
    line_spacing=1.15,
    font=FONT,
    page_break_before=False,
):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if page_break_before:
        pf.page_break_before = True

    if text:
        run = p.add_run(text)
        set_run_font(run, size_pt=size, bold=bold, italic=italic, font=font)
    return p


def add_mixed_para(doc, parts, *, align="justify", size=12, space_before=0, space_after=10, line_spacing=1.15):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    for text, bold in parts:
        run = p.add_run(text)
        set_run_font(run, size_pt=size, bold=bold)
    return p


def add_code_block(doc, code: str, size=9):
    for line in code.splitlines():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.4)
        run = p.add_run(line if line else " ")
        set_run_font(run, size_pt=size, bold=False, font=FONT_MONO)


def section_heading(doc, text):
    return add_para(
        doc,
        text,
        size=14,
        bold=True,
        align="center",
        space_after=14,
        line_spacing=1.15,
        page_break_before=True,
    )


def build():
    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ---------- COVER ----------
    add_para(doc, "", space_after=12)
    add_para(doc, "Blockchain Mini Project Report", size=16, bold=True, align="center", space_after=16, line_spacing=1.0)
    add_para(doc, "Decentralized Voting System", size=22, bold=True, align="center", space_before=4, space_after=16, line_spacing=1.15)
    add_para(doc, "Submitted fulfilment of the requirements of the", size=12, align="center", space_after=0, line_spacing=1.15)
    add_para(doc, "Blockchain Mini Project Semester VII in degree of", size=12, align="center", space_after=0, line_spacing=1.15)
    add_para(doc, "Bachelor of Computer Engineering", size=12, align="center", space_after=12, line_spacing=1.15)
    add_para(doc, "BY", size=12, align="center", space_after=6, line_spacing=1.0)
    for name in ["Riya Jha (47)", "Rugved Khandake (60)"]:
        add_para(doc, name, size=13, bold=True, align="center", space_after=2, line_spacing=1.15)
    add_para(doc, "", space_after=8)
    add_para(doc, "Guide", size=12, align="center", space_after=2, line_spacing=1.0)
    add_para(doc, "Asst. Prof. Manasi Churi", size=12, align="center", space_after=12, line_spacing=1.15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run()
    run.add_picture(str(FIGS / "college-logo.png"), width=Inches(1.55))

    add_para(doc, "Department Of Computer Engineering", size=13, bold=True, align="center", space_after=2, line_spacing=1.15)
    add_para(doc, "Shree L. R. Tiwari College Of Engineering", size=13, bold=True, align="center", space_after=2, line_spacing=1.15)
    add_para(doc, "Kanakia Park, Mira Road (E), Thane -401 107, Maharashtra.", size=12, bold=True, align="center", space_after=12, line_spacing=1.15)
    add_para(doc, "2026-27", size=14, bold=True, align="center", space_after=0, line_spacing=1.0)

    # ---------- APPROVAL ----------
    add_para(doc, "Blockchain Mini Project Report Approval", size=16, bold=True, align="center", space_after=28, line_spacing=1.15, page_break_before=True)
    add_mixed_para(
        doc,
        [
            ("This project report entitled ", False),
            ("\u201cDecentralized Voting System\u201d", True),
            (" by ", False),
            ("Ms. Riya Jha (47) and Mr. Rugved Khandake (60)", True),
            (
                " is approved for the Mini Project Semester VII project work of B.E CMPN at Shree L.R.Tiwari College of Engineering, in the academic year 2026-2027.",
                False,
            ),
        ],
        align="justify", size=12, space_after=6, line_spacing=1.5,
    )
    for _ in range(10):
        add_para(doc, "", space_after=12)
    add_para(doc, "Internal Examiner", size=12, bold=True, align="right", space_after=0)

    # ---------- DECLARATION ----------
    add_para(doc, "Declaration", size=16, bold=True, align="center", space_after=20, line_spacing=1.15, page_break_before=True)
    decl = (
        "I declare that this written submission represents my ideas in my own words and where "
        "others' ideas or words have been included, I have adequately cited and referenced the "
        "original sources. I also declare that I have adhered to all principles of academic "
        "honesty and integrity and have not misrepresented or fabricated or falsified any "
        "idea/data/fact/source in my submission. I understand that any violation of the above "
        "will be cause for disciplinary action by the Institute and can also evoke penal action "
        "from the sources which have thus not been properly cited or from whom proper permission "
        "has not been taken when needed."
    )
    add_para(doc, decl, size=12, align="justify", space_after=36, line_spacing=1.5)
    for name in ["Ms. Riya Jha (47)", "Mr. Rugved Khandake (60)"]:
        add_para(doc, "________________________", size=12, align="right", space_before=18, space_after=0, line_spacing=1.0)
        add_para(doc, "Signature", size=11, align="right", space_after=0, line_spacing=1.0)
        add_para(doc, name, size=12, bold=True, align="right", space_after=6, line_spacing=1.15)
    add_para(doc, "", space_after=36)
    add_para(doc, "Date:", size=12, align="left", space_after=6)
    add_para(doc, "Place:", size=12, align="left", space_after=0)

    # ---------- TOC ----------
    add_para(doc, "Table of Contents", size=14, bold=True, align="left", space_after=14, line_spacing=1.15, page_break_before=True)
    for item in [
        "1. Introduction",
        "2. System Architecture",
        "3. Implementation",
        "4. Results",
        "5. Conclusion",
        "6. References",
    ]:
        add_para(doc, item, size=12, align="left", space_after=4, line_spacing=1.5)

    # ---------- 1. INTRODUCTION ----------
    section_heading(doc, "1. Introduction")
    for t in [
        (
            "Traditional voting systems often rely on centralized authorities for voter rolls, "
            "ballot counting, and result publication. Centralization can raise concerns about "
            "transparency, single points of failure, and the difficulty of independently verifying "
            "that each eligible voter cast exactly one ballot and that tallies were not altered "
            "after the fact."
        ),
        (
            "Blockchain technology provides an immutable, shared ledger and programmable rules "
            "via smart contracts. An election contract can record candidate lists, prevent "
            "double voting with a hasVoted mapping, accept votes from any address while the "
            "poll is open, and publish a winner only after the owner ends the election. Events "
            "such as VoteCast and ElectionEnded give an auditable trail of participation."
        ),
        (
            "This mini-project implements a Decentralized Voting System as a local MVP for the "
            "Blockchain laboratory (Semester VII). The contract owner creates an election with a "
            "title and candidate names. Any wallet may vote once for one candidate. The owner "
            "ends the election, after which the winner is readable on-chain. The stack uses "
            "Hardhat for local development, Solidity for on-chain logic, and a lightweight "
            "HTML/CSS/JavaScript frontend with ethers.js v6."
        ),
    ]:
        add_para(doc, t, size=12, align="justify", space_after=10, line_spacing=1.5)

    add_para(doc, "Objectives", size=12, bold=True, align="left", space_before=6, space_after=6)
    for i, o in enumerate([
        "Design a Solidity smart contract for creating elections, casting one vote per address, and ending the poll.",
        "Emit VoteCast and ElectionEnded events and expose candidate tallies and getWinner after end.",
        "Provide unit tests and a runnable local demo (Hardhat node + MetaMask-compatible UI).",
        "Document architecture, implementation, and results for academic evaluation.",
    ], 1):
        add_para(doc, f"{i}. {o}", size=12, align="justify", space_after=4, line_spacing=1.35)

    add_para(doc, "Scope (MVP)", size=12, bold=True, align="left", space_before=10, space_after=6)
    for s in [
        "Local Hardhat network only; no public-chain deployment.",
        "One election per contract deployment; no voter identity KYC beyond wallet address.",
        "No token-weighted or quadratic voting; one address, one vote.",
    ]:
        add_para(doc, f"\u2022  {s}", size=12, align="justify", space_after=4, line_spacing=1.35)

    # ---------- 2. SYSTEM ARCHITECTURE ----------
    section_heading(doc, "2. System Architecture")
    add_para(
        doc,
        (
            "The system involves two primary actors. The owner (contract deployer) creates an "
            "election with a title and at least two candidate names, and may later end the "
            "election. A voter connects a wallet, selects one candidate while the election is "
            "open, and submits a vote transaction. The contract rejects a second vote from the "
            "same address and rejects votes after the election has ended."
        ),
        size=12, align="justify", space_after=10, line_spacing=1.5,
    )
    add_para(doc, "The system is organized in three layers:", size=12, align="justify", space_after=6, line_spacing=1.5)
    for bold_part, rest in [
        ("Smart contract layer \u2014 ", "VotingSystem.sol stores election metadata, Candidate records, hasVoted mapping, and emits ElectionCreated / VoteCast / ElectionEnded."),
        ("Tooling layer \u2014 ", "Hardhat compiles, tests, and deploys the contract. The deploy script seeds a sample election and writes the address and ABI for the UI."),
        ("Presentation layer \u2014 ", "A vanilla single-page app loads the ABI and address, connects a wallet (or runs read-only against the local node), lists candidates with tallies, submits votes, and lets the owner end the election."),
    ]:
        add_mixed_para(doc, [(f"\u2022  {bold_part}", True), (rest, False)], align="justify", size=12, space_after=6, line_spacing=1.4)

    add_para(doc, "Logical architecture (local MVP):", size=12, bold=True, align="left", space_before=8, space_after=6)
    arch = """+---------------------+     JSON-RPC      +--------------------------+
|  Frontend (HTML/JS) | <---------------> |  Hardhat Local Node      |
|  ethers.js (CDN)    |                   |  VotingSystem.sol        |
+---------------------+                   +--------------------------+
         |                                           |
         |  contract-address.json + abi.json         |
         +------------ written by deploy.js ---------+"""
    add_code_block(doc, arch, size=8)

    add_para(doc, "On-chain data and key flows", size=12, bold=True, align="left", space_before=12, space_after=6)
    add_para(
        doc,
        (
            "Each Candidate struct holds id, name, and voteCount. createElection (owner-only) "
            "sets electionTitle, marks electionCreated, and stores candidates. vote checks that "
            "the election is open, the sender has not voted, and the candidate id is valid, then "
            "sets hasVoted[msg.sender] and increments the tally. endElection (owner-only) sets "
            "electionEnded and emits ElectionEnded with the winning candidate. getWinner is "
            "available only after end; ties prefer the lowest candidate id among leaders."
        ),
        size=12, align="justify", space_after=6, line_spacing=1.5,
    )

    # ---------- 3. IMPLEMENTATION ----------
    section_heading(doc, "3. Implementation")
    add_para(
        doc,
        (
            "The core contract is contracts/VotingSystem.sol (Solidity ^0.8.20). Custom errors "
            "provide clear reverts (NotOwner, AlreadyVoted, ElectionAlreadyEnded, InvalidCandidate, "
            "and others). View helpers include getCandidate, getAllCandidates, candidateCount, "
            "hasVoted, and getWinner (post-end)."
        ),
        size=12, align="justify", space_after=8, line_spacing=1.4,
    )
    add_para(doc, "Key Solidity functions:", size=12, bold=True, align="left", space_after=4)

    code_snippet = r'''function createElection(string calldata title, string[] calldata names)
    external onlyOwner
{
    if (electionCreated) revert ElectionAlreadyCreated();
    if (bytes(title).length == 0) revert EmptyTitle();
    if (names.length < 2) revert NeedCandidates();
    electionTitle = title;
    electionCreated = true;
    for (uint256 i = 0; i < names.length; i++) {
        candidateCount++;
        _candidates[candidateCount] = Candidate({
            id: candidateCount, name: names[i], voteCount: 0
        });
    }
    emit ElectionCreated(title, candidateCount);
}

function vote(uint256 candidateId) external {
    if (!electionCreated) revert ElectionNotCreated();
    if (electionEnded) revert ElectionAlreadyEnded();
    if (hasVoted[msg.sender]) revert AlreadyVoted();
    if (candidateId == 0 || candidateId > candidateCount)
        revert InvalidCandidate();
    hasVoted[msg.sender] = true;
    _candidates[candidateId].voteCount += 1;
    emit VoteCast(msg.sender, candidateId);
}

function endElection() external onlyOwner {
    if (!electionCreated) revert ElectionNotCreated();
    if (electionEnded) revert ElectionAlreadyEnded();
    electionEnded = true;
    (uint256 winnerId, string memory winnerName, uint256 winnerVotes)
        = _computeWinner();
    emit ElectionEnded(winnerId, winnerName, winnerVotes);
}'''
    add_code_block(doc, code_snippet, size=8)

    add_para(
        doc,
        (
            "Hardhat project layout includes hardhat.config.js (Solidity 0.8.20, localhost), "
            "scripts/deploy.js (deploys, seeds a sample election with three candidates, writes "
            "frontend/contract-address.json and frontend/abi.json), scripts/smoke.js (vote + end + "
            "winner), and test/VotingSystem.js (Mocha/Chai coverage of create, vote, double-vote "
            "revert, end, winner, cannot vote after end, non-owner create, and getWinner before "
            "end). The frontend under frontend/ uses index.html, styles.css, and app.js with "
            "ethers.js v6 from CDN."
        ),
        size=12, align="justify", space_before=8, space_after=6, line_spacing=1.4,
    )

    # ---------- TECHNOLOGIES USED ----------
    section_heading(doc, "Technologies Used")
    for name, role in [
        ("Solidity ^0.8.20", "Smart contract language for VotingSystem.sol."),
        ("Hardhat", "Compile, unit test, local blockchain node, and deployment."),
        ("ethers.js v6", "Contract interaction in tests and the browser frontend."),
        ("Node.js", "Runtime for Hardhat tooling and npm scripts."),
        ("Vanilla HTML / CSS / JavaScript", "Voting UI without a heavy framework."),
        ("MetaMask (optional)", "Injected wallet for interactive vote and owner demos."),
    ]:
        add_mixed_para(
            doc, [(f"\u2022  {name} \u2014 ", True), (role, False)],
            align="justify", size=12, space_after=8, line_spacing=1.4,
        )

    # ---------- 4. RESULTS ----------
    section_heading(doc, "4. Results")
    add_para(
        doc,
        (
            "The project was verified on a local Hardhat environment. Compilation succeeds with "
            "npx hardhat compile. All eight unit tests pass under npx hardhat test (create election, "
            "vote, double-vote revert, end election, winner, cannot vote after end, non-owner "
            "create, and getWinner before end). A smoke run with a live Hardhat node successfully "
            "casts a vote, ends the election, and confirms the expected winner."
        ),
        size=11, align="justify", space_after=4, line_spacing=1.2,
    )
    add_para(
        doc,
        (
            "After deployment, the static frontend shows the seeded election title and candidates "
            "with live vote counts on the local network. Connecting a Hardhat account allows "
            "voting while the poll is open. The owner account can end the election; the UI then "
            "displays the winner. Read-only JsonRpcProvider mode works without MetaMask for "
            "demonstration of tallies."
        ),
        size=11, align="justify", space_after=4, line_spacing=1.2,
    )

    figures = [
        ("fig1-election-open.png", "Fig. 1: Open election with candidates and live tallies"),
        ("fig2-election-ended.png", "Fig. 2: Election ended \u2014 winner displayed"),
        ("fig3-demo-help.png", "Fig. 3: Demo help panel (Hardhat steps)"),
    ]

    # Stack screenshots vertically so they print large enough to read.
    for fig_path, caption in figures:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(str(FIGS / fig_path), width=Inches(5.8))

        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(4)
        cr = cp.add_run(caption)
        cr.italic = True
        cr.font.size = Pt(10)
        cr.font.name = "Liberation Serif"

    add_para(
        doc,
        (
            "Summary: 8/8 Hardhat tests passed; smoke vote + end + winner succeeded; the screenshots "
            "above confirm the local voting demo path used for laboratory evaluation."
        ),
        size=11, align="justify", space_before=6, space_after=4, line_spacing=1.25,
    )

    # ---------- 5. CONCLUSION ----------
    section_heading(doc, "5. Conclusion")
    add_para(
        doc,
        (
            "This mini-project delivered a working Decentralized Voting System MVP suitable for "
            "a Blockchain laboratory demonstration. A Solidity contract manages election creation, "
            "one-vote-per-address casting, election closure, and winner publication. Hardhat "
            "provides compilation, automated tests (8/8 passing), and a local chain. A simple web "
            "UI lets users interact through a wallet or in read-only mode against the deployed "
            "contract address."
        ),
        size=12, align="justify", space_after=10, line_spacing=1.5,
    )
    add_para(
        doc,
        (
            "The design deliberately stays small: no public-chain deployment, no off-chain voter "
            "registry, and a single election per deployment. Future work may include multiple "
            "concurrent elections, commit\u2013reveal ballots for privacy, Merkle-proof voter "
            "eligibility, or time-locked voting windows. Within the stated scope, the happy "
            "path\u2014from create to vote to end to winner\u2014is implemented, tested, and documented."
        ),
        size=12, align="justify", space_after=6, line_spacing=1.5,
    )

    # ---------- 6. REFERENCES ----------
    section_heading(doc, "6. References")
    refs = [
        "Ethereum Foundation, \u201cEthereum Developer Documentation,\u201d https://ethereum.org/developers/docs/ (accessed 2026).",
        "Solidity Team, \u201cSolidity Documentation,\u201d https://docs.soliditylang.org/ (accessed 2026).",
        "Nomic Foundation, \u201cHardhat Documentation,\u201d https://hardhat.org/docs (accessed 2026).",
        "ethers.js, \u201cethers.js v6 Documentation,\u201d https://docs.ethers.org/v6/ (accessed 2026).",
        "Ethereum Foundation, \u201cSmart Contract Security Best Practices \u2014 Access Control,\u201d applied via onlyOwner for create/end.",
        "Shree L. R. Tiwari College Of Engineering, Department of Computer Engineering \u2014 Blockchain Mini Project guidelines (internal).",
    ]
    for i, r in enumerate(refs, 1):
        add_para(doc, f"{i}. {r}", size=12, align="justify", space_after=8, line_spacing=1.35)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))
    print(f"Wrote {OUT_DOCX}")


if __name__ == "__main__":
    build()
